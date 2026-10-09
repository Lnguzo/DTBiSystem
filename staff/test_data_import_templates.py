from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from .data_services import IMPORT_FIELDS, IMPORT_MODELS, _value_for_field, guess_field_mapping
from .models import UserProfile


class DataImportTemplateTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(
            username='importstaff', password='secret1234', is_superuser=True,
        )
        UserProfile.objects.get_or_create(user=self.user)
        self.client.force_login(self.user)

    def test_import_page_renders_upload_controls_without_template_links(self):
        response = self.client.get(reverse('staff:data_import'))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'name="dataset"')
        self.assertContains(response, 'name="file"')
        self.assertNotContains(response, 'Download Startups template')

    def test_startup_template_matches_startup_record_fields_and_auto_maps(self):
        response = self.client.get(
            reverse('staff:data_import_template', kwargs={'dataset': 'startup'})
        )

        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content), read_only=True)
        try:
            headers = list(next(workbook.active.iter_rows(values_only=True)))
        finally:
            workbook.close()

        expected_headers = [
            IMPORT_MODELS['startup']._meta.get_field(field).verbose_name.title()
            for field in IMPORT_FIELDS['startup']
        ]
        self.assertEqual(headers, expected_headers)
        mapping = guess_field_mapping('startup', headers)
        self.assertEqual(set(mapping), set(IMPORT_FIELDS['startup']))
        self.assertIn('contact_person', mapping)
        self.assertIn('contact_address', mapping)

    def test_startup_choice_labels_convert_to_platform_values(self):
        self.assertEqual(_value_for_field('Individual', 'startup_type', 'startup'), 'individual')
        self.assertEqual(_value_for_field('Active', 'status', 'startup'), 'active')
        self.assertEqual(_value_for_field('Draft', 'contract_status', 'startup'), 'draft')
