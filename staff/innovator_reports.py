"""Public Innovator directory exports.

Exports include only fields displayed in published Innovator profiles. In
particular, private phone numbers and staff-only programme notes are omitted.
"""
from html import escape
import io
import re

from django.db.models import Q
from django.http import Http404, HttpResponse
from django.utils import timezone
from django.views.decorators.http import require_GET

from .models import InnovatorProfile


def public_innovator_queryset(params):
    """Apply the same published, search, and catalogue filters as the directory."""
    profiles = InnovatorProfile.objects.filter(is_published=True)
    catalogue = params.get('catalogue', '').strip()
    if catalogue in {'agritech', 'makisatu'}:
        profiles = profiles.filter(catalogue=catalogue)
    query = params.get('q', '').strip()
    if query:
        profiles = profiles.filter(
            Q(innovator_name__icontains=query)
            | Q(project_name__icontains=query)
            | Q(description__icontains=query)
        )
    return profiles.order_by('catalogue', 'sort_order', 'project_name')


def _public_rows(profiles):
    return [
        [
            profile.project_name,
            profile.innovator_name,
            profile.get_catalogue_display(),
            profile.description,
            profile.contact_email,
            profile.website,
            profile.source_url,
        ]
        for profile in profiles
    ]


def _clean(value, limit=None):
    text = re.sub(r'\s+', ' ', str(value or '')).strip()
    if limit and len(text) > limit:
        text = text[:limit - 3].rstrip() + '...'
    return text


def _excel_report(profiles, params):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    rows = _public_rows(profiles)
    counts = {'agritech': 0, 'makisatu': 0}
    for profile in profiles:
        if profile.catalogue in counts:
            counts[profile.catalogue] += 1

    book = Workbook()
    book.properties.title = 'DTBi Innovator Directory Summary'
    summary = book.active
    summary.title = 'Summary'
    summary.append(['DTBi | BUNI Innovator Directory', ''])
    summary.merge_cells('A1:B1')
    summary['A1'].font = Font(bold=True, size=16, color='FFFFFF')
    summary['A1'].fill = PatternFill('solid', fgColor='173C78')
    summary.append(['Report generated', timezone.localtime().strftime('%d %b %Y, %H:%M %Z')])
    summary.append(['Published projects included', len(rows)])
    summary.append(['AgriTech and Food Systems', counts['agritech']])
    summary.append(['MAKISATU Innovation Catalogue', counts['makisatu']])
    summary.append(['Search filter', params.get('q', '').strip() or 'All projects'])
    summary.append(['Catalogue filter', params.get('catalogue', '').strip() or 'All catalogues'])
    summary.append(['Privacy note', 'Public profile details only. Phone numbers and staff-only programme notes are excluded.'])
    summary.column_dimensions['A'].width = 34
    summary.column_dimensions['B'].width = 92
    for row in summary.iter_rows(min_row=2):
        row[0].font = Font(bold=True, color='173C78')
        row[1].alignment = Alignment(vertical='top', wrap_text=True)
    summary.freeze_panes = 'A2'

    directory = book.create_sheet('Published Innovators')
    headers = ['Project', 'Innovator / group', 'Catalogue', 'Description', 'Public email', 'Website', 'Public source']
    directory.append(headers)
    for cell in directory[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='173C78')
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    for row in rows:
        directory.append([_clean(value) for value in row])
    widths = [30, 26, 34, 76, 32, 38, 42]
    for index, width in enumerate(widths, start=1):
        directory.column_dimensions[get_column_letter(index)].width = width
    for row in directory.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    directory.freeze_panes = 'A2'
    directory.auto_filter.ref = directory.dimensions

    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def _pdf_report(profiles, params):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    rows = _public_rows(profiles)
    output = io.BytesIO()
    page_size = landscape(A4)
    doc = SimpleDocTemplate(
        output, pagesize=page_size, title='DTBi | BUNI Innovator Directory Summary',
        leftMargin=28, rightMargin=28, topMargin=34, bottomMargin=38,
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle('InnovatorReportTitle', parent=styles['Title'], textColor=colors.HexColor('#173c78'), spaceAfter=7)
    note_style = ParagraphStyle('InnovatorReportNote', parent=styles['BodyText'], fontSize=8, leading=11, textColor=colors.HexColor('#53657a'))
    cell_style = ParagraphStyle('InnovatorReportCell', parent=styles['BodyText'], fontSize=7.1, leading=9, spaceAfter=0)
    head_style = ParagraphStyle('InnovatorReportHead', parent=cell_style, textColor=colors.white, fontName='Helvetica-Bold')

    def cell(value, style):
        safe = escape(_clean(value, 420) or '-')
        safe = safe.encode('latin-1', 'replace').decode('latin-1')
        return Paragraph(safe, style)

    story = [
        Paragraph('DTBi | BUNI Innovator Directory', title_style),
        Paragraph(
            f"{len(rows)} published projects · Search: {escape(_clean(params.get('q')) or 'all')} · "
            f"Catalogue: {escape(_clean(params.get('catalogue')) or 'all')}", note_style,
        ),
        Paragraph('Public profile details only. Confidential phone numbers and staff-only programme notes are excluded.', note_style),
        Spacer(1, 12),
    ]
    headers = ['Project', 'Innovator / group', 'Catalogue', 'Description', 'Public email', 'Website', 'Public source']
    usable_width = page_size[0] - 56
    weights = [1.35, 1.15, 1.5, 3.45, 1.35, 1.35, 1.4]
    col_widths = [usable_width * weight / sum(weights) for weight in weights]
    data = [[cell(header, head_style) for header in headers]]
    data.extend([[cell(value, cell_style) for value in row] for row in rows])
    table = Table(data, colWidths=col_widths, repeatRows=1, hAlign='LEFT')
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#173c78')),
        ('GRID', (0, 0), (-1, -1), .35, colors.HexColor('#d6deea')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f3f6fb')]),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('LEFTPADDING', (0, 0), (-1, -1), 4),
        ('RIGHTPADDING', (0, 0), (-1, -1), 4),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
    ]))
    story.append(table)

    def footer(canvas, _doc):
        canvas.saveState()
        canvas.setFont('Helvetica', 8)
        canvas.setFillColor(colors.HexColor('#64748b'))
        canvas.drawString(28, 20, 'DTBi | BUNI · Public Innovator directory summary')
        canvas.drawRightString(page_size[0] - 28, 20, f'Page {_doc.page}')
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()


@require_GET
def innovators_export(request, file_format):
    if file_format not in {'xlsx', 'pdf'}:
        raise Http404
    profiles = list(public_innovator_queryset(request.GET))
    if file_format == 'xlsx':
        payload = _excel_report(profiles, request.GET)
        content_type = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    else:
        payload = _pdf_report(profiles, request.GET)
        content_type = 'application/pdf'
    response = HttpResponse(payload, content_type=content_type)
    response['Content-Disposition'] = f'attachment; filename="dtbi-innovators-summary.{file_format}"'
    response['X-Report-Rows'] = str(len(profiles))
    return response
